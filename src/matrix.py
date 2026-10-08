"""Build a literature-review matrix from extracted papers for CSV and Excel export."""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

COLUMNS = [
    "Paper",
    "Title",
    "Year",
    "Authors",
    "Problem / Task",
    "Dataset",
    "Dataset URL",
    "Features",
    "Model / Method",
    "Metrics",
    "Limitations",
    "Paper URL",
]

NOTICE = (
    "Use this to decide what to read. Open the paper, then write your own notes. "
    "Do not paste this into a proposal. Dataset links are from the papers plus public search pages."
)

COLUMN_WIDTHS = {
    "Paper": 10,
    "Title": 46,
    "Year": 8,
    "Authors": 28,
    "Problem / Task": 38,
    "Dataset": 24,
    "Dataset URL": 36,
    "Features": 24,
    "Model / Method": 28,
    "Metrics": 24,
    "Limitations": 38,
    "Paper URL": 36,
}

DATASET_COLUMNS = [
    "Dataset",
    "Kind",
    "Size",
    "Access",
    "Strengths",
    "Limitations",
    "Feasibility",
    "About",
    "Why useful",
    "Papers",
    "Homepage",
    "Hugging Face",
    "Web search",
]

DATASET_WIDTHS = {
    "Dataset": 28,
    "Kind": 14,
    "Size": 16,
    "Access": 14,
    "Strengths": 32,
    "Limitations": 32,
    "Feasibility": 22,
    "About": 36,
    "Why useful": 36,
    "Papers": 28,
    "Homepage": 40,
    "Hugging Face": 40,
    "Web search": 40,
}

_HEADER_FONT = Font(bold=True, color="1C1915", name="Calibri", size=11)
_HEADER_FILL = PatternFill("solid", fgColor="EAD9B2")
_NOTICE_FILL = PatternFill("solid", fgColor="F4EFE4")
_WRAP = Alignment(wrap_text=True, vertical="top")
_HEADER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="left")
_THIN = Border(
    left=Side(style="thin", color="DDD4C4"),
    right=Side(style="thin", color="DDD4C4"),
    top=Side(style="thin", color="DDD4C4"),
    bottom=Side(style="thin", color="DDD4C4"),
)
_LINK_FONT = Font(name="Calibri", size=10, color="0563C1", underline="single")
_BODY_FONT = Font(name="Calibri", size=10)


def _clean(value: Any, limit: int = 2000) -> str:
    text = " ".join(str(value or "").split())
    if len(text) <= limit:
        return text
    return text[: limit - 1].rstrip() + "..."


def _authors(paper: dict) -> str:
    names = paper.get("authors") or []
    if not names:
        return ""
    shown = names[:4]
    extra = f" +{len(names) - 4}" if len(names) > 4 else ""
    return ", ".join(shown) + extra


def _best_dataset_url(item: dict) -> str:
    return str(item.get("homepage") or item.get("hf_url") or item.get("search_url") or "").strip()


def _match_datasets(paper: dict, datasets: list[dict]) -> list[dict]:
    blob = " ".join(
        [
            str(paper.get("datasets") or ""),
            str(paper.get("title") or ""),
        ]
    ).lower()
    hits: list[dict] = []
    for item in datasets or []:
        name = str(item.get("name") or "").strip()
        if len(name) < 3:
            continue
        if name.lower() in blob:
            hits.append(item)
    if hits:
        return hits
    rank = paper.get("rank")
    for item in datasets or []:
        for linked in item.get("papers") or []:
            if linked.get("rank") == rank:
                hits.append(item)
                break
    return hits


def matrix_rows(papers: list[dict], datasets: list[dict] | None = None) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    catalog = datasets or []
    for index, paper in enumerate(papers or [], start=1):
        rank = paper.get("rank") or index
        metrics = paper.get("metrics") or paper.get("key_findings") or ""
        features = paper.get("features") or ""
        matched = _match_datasets(paper, catalog)
        names = ", ".join(item.get("name") or "" for item in matched if item.get("name"))
        urls = [url for url in (_best_dataset_url(item) for item in matched) if url]
        rows.append(
            {
                "Paper": f"P{rank}",
                "Title": _clean(paper.get("title"), 240),
                "Year": str(paper.get("year") or ""),
                "Authors": _authors(paper),
                "Problem / Task": _clean(paper.get("problem")),
                "Dataset": names or _clean(paper.get("datasets")),
                "Dataset URL": " | ".join(urls),
                "Features": _clean(features),
                "Model / Method": _clean(paper.get("method")),
                "Metrics": _clean(metrics),
                "Limitations": _clean(paper.get("limitations")),
                "Paper URL": str(paper.get("url") or ""),
            }
        )
    return rows


def dataset_rows(datasets: list[dict]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for item in datasets or []:
        papers = item.get("papers") or []
        paper_labels = ", ".join(
            f"P{p.get('rank')}" for p in papers if p.get("rank") is not None
        )
        rows.append(
            {
                "Dataset": _clean(item.get("name"), 120),
                "Kind": _clean(item.get("kind"), 40),
                "Size": _clean(item.get("size"), 80),
                "Access": _clean(item.get("access"), 40),
                "Strengths": _clean(item.get("strengths") or item.get("about"), 400),
                "Limitations": _clean(item.get("limitations"), 400),
                "Feasibility": _clean(
                    " ".join(
                        part
                        for part in (
                            item.get("feasibility") or "",
                            item.get("feasibility_note") or "",
                        )
                        if part
                    ),
                    80,
                ),
                "About": _clean(item.get("about"), 400),
                "Why useful": _clean(item.get("why_useful"), 400),
                "Papers": paper_labels,
                "Homepage": str(item.get("homepage") or ""),
                "Hugging Face": str(item.get("hf_url") or ""),
                "Web search": str(item.get("search_url") or ""),
            }
        )
    return rows


def matrix_csv(papers: list[dict], datasets: list[dict] | None = None) -> str:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=COLUMNS, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(matrix_rows(papers, datasets))
    return buffer.getvalue()


def matrix_xlsx(papers: list[dict], datasets: list[dict] | None = None) -> bytes:
    wb = Workbook()
    _fill_sheet(
        wb.active,
        title="Literature matrix",
        columns=COLUMNS,
        widths=COLUMN_WIDTHS,
        rows=matrix_rows(papers, datasets),
        link_columns={"Dataset URL", "Paper URL"},
    )
    if datasets:
        _fill_sheet(
            wb.create_sheet("Datasets"),
            title="Datasets",
            columns=DATASET_COLUMNS,
            widths=DATASET_WIDTHS,
            rows=dataset_rows(datasets),
            link_columns={"Homepage", "Hugging Face", "Web search"},
        )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _fill_sheet(
    ws: Worksheet,
    *,
    title: str,
    columns: list[str],
    widths: dict[str, float],
    rows: list[dict[str, str]],
    link_columns: set[str],
) -> None:
    ws.title = title[:31]
    last_col = get_column_letter(len(columns))

    for col_index, name in enumerate(columns, start=1):
        cell = ws.cell(1, col_index, name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN
        ws.column_dimensions[get_column_letter(col_index)].width = widths.get(name, 24)

    ws.row_dimensions[1].height = 32

    for row_index, row in enumerate(rows, start=2):
        longest = 1
        for col_index, name in enumerate(columns, start=1):
            value = row.get(name) or ""
            cell = ws.cell(row_index, col_index)
            cell.alignment = _WRAP
            cell.border = _THIN
            if name in link_columns and value.startswith("http"):
                _set_links(cell, value)
            else:
                cell.value = value
                cell.font = _BODY_FONT
            longest = max(longest, _wrapped_lines(value, widths.get(name, 24)))
        ws.row_dimensions[row_index].height = min(120, max(22, longest * 15))

    data_last = max(2, 1 + len(rows))
    if not rows:
        empty = ws.cell(2, 1, "No datasets named in these papers.")
        empty.font = _BODY_FONT
        data_last = 2

    notice_row = data_last + 2
    notice_cell = ws.cell(notice_row, 1, NOTICE)
    notice_cell.font = Font(name="Calibri", size=10, italic=True, color="6F675C")
    notice_cell.fill = _NOTICE_FILL
    notice_cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells(start_row=notice_row, start_column=1, end_row=notice_row, end_column=len(columns))
    ws.row_dimensions[notice_row].height = 36

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{last_col}{data_last}"
    ws.sheet_view.showGridLines = False


def _set_links(cell: Cell, value: str) -> None:
    parts = [part.strip() for part in value.split("|") if part.strip().startswith("http")]
    if not parts:
        cell.value = value
        cell.font = _BODY_FONT
        return
    first = parts[0]
    cell.value = first
    cell.hyperlink = first
    cell.font = _LINK_FONT
    if len(parts) > 1:
        cell.value = first + f" (+{len(parts) - 1} more)"


def _wrapped_lines(text: str, width: float) -> int:
    if not text:
        return 1
    chars = max(8, int(width))
    return max(1, (len(text) + chars - 1) // chars)


def safe_filename(idea: str, task_id: str, ext: str = "xlsx") -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", (idea or "research")[:40]).strip("-").lower()
    short = (task_id or "export")[:8]
    return f"literature-matrix-{slug or 'research'}-{short}.{ext}"


GAP_CHECK_COLUMNS = [
    "Gap #",
    "Research Gap Title",
    "Status",
    "Originating Paper(s)",
    "Originating Paper Year",
    "Candidates Checked",
    "Solving Paper",
    "Solving Year & Venue",
    "Solving Paper URL",
    "Downstream Gap (Opportunity)",
    "Downstream Status",
    "AI Verdict Reasoning",
    "Search Queries Used",
    "Gap Details",
]

GAP_CHECK_WIDTHS = {
    "Gap #": 8,
    "Research Gap Title": 34,
    "Status": 15,
    "Originating Paper(s)": 32,
    "Originating Paper Year": 16,
    "Candidates Checked": 14,
    "Solving Paper": 36,
    "Solving Year & Venue": 18,
    "Solving Paper URL": 36,
    "Downstream Gap (Opportunity)": 38,
    "Downstream Status": 16,
    "AI Verdict Reasoning": 45,
    "Search Queries Used": 34,
    "Gap Details": 40,
}


def gap_check_rows(
    gc_results: list[dict],
    all_gaps: list[dict] | None = None,
    papers: list[dict] | None = None,
) -> list[dict[str, str]]:
    rows = []
    for idx, r in enumerate(gc_results):
        gap_title = str(r.get("gap_title") or f"Gap {idx+1}")
        solved = bool(r.get("solved"))
        status = "SOLVED" if solved else "STILL OPEN"

        # Originating paper(s) resolution
        supporters = r.get("supported_by") or []
        if not supporters and all_gaps:
            match = next(
                (g for g in all_gaps if (g.get("title") or "").strip().lower() == gap_title.strip().lower()),
                None,
            )
            if match and match.get("supported_by"):
                supporters = match.get("supported_by") or []

        paper_labels = []
        paper_years = []
        for s in supporters:
            if isinstance(s, dict):
                pid = s.get("id") or (f"P{s.get('rank')}" if s.get("rank") else "Paper")
                title = s.get("title") or ""
                yr = str(s.get("year") or s.get("date") or "")
                if not yr and papers:
                    match_p = next(
                        (
                            p for p in papers
                            if (p.get("id") and p.get("id") == pid)
                            or (p.get("rank") and f"P{p.get('rank')}" == pid)
                            or (p.get("title") and title and p.get("title").strip().lower() == title.strip().lower())
                        ),
                        None,
                    )
                    if match_p:
                        yr = str(match_p.get("year") or match_p.get("date") or "")
                label = f"{pid}: {title}" if title else pid
                if yr:
                    label += f" ({yr[:4]})"
                    paper_years.append(yr[:4])
                paper_labels.append(label)
            else:
                paper_labels.append(str(s))

        solving_papers = r.get("solving_papers") or []
        sp = solving_papers[0] if solving_papers else {}
        sp_title = sp.get("title") or ("None (Open Gap)" if not solved else "")
        sp_meta = f"{sp.get('year', '')} [{sp.get('source', '')}]".strip() if sp else ""
        sp_url = sp.get("url") or ""

        # Downstream info
        ds = r.get("downstream") or {}
        ds_gap = ds.get("gap") or ""
        ds_status = ""
        if ds:
            ds_status = "SOLVED" if ds.get("solved") else "STILL OPEN"

        queries = r.get("queries_used") or []
        queries_str = " | ".join(queries) if isinstance(queries, list) else str(queries)

        candidates_count = r.get("candidates_found") or len(r.get("evaluated_papers") or [])

        rows.append({
            "Gap #": f"Gap {idx+1}",
            "Research Gap Title": gap_title,
            "Status": status,
            "Originating Paper(s)": "\n".join(paper_labels) if paper_labels else "N/A",
            "Originating Paper Year": ", ".join(dict.fromkeys(paper_years)) if paper_years else "N/A",
            "Candidates Checked": str(candidates_count),
            "Solving Paper": sp_title,
            "Solving Year & Venue": sp_meta,
            "Solving Paper URL": sp_url,
            "Downstream Gap (Opportunity)": ds_gap,
            "Downstream Status": ds_status,
            "AI Verdict Reasoning": str(r.get("reasoning") or ""),
            "Search Queries Used": queries_str,
            "Gap Details": str(r.get("gap_detail") or ""),
        })
    return rows


def gap_check_xlsx(
    gc_results: list[dict],
    all_gaps: list[dict] | None = None,
    papers: list[dict] | None = None,
) -> bytes:
    wb = Workbook()
    _fill_sheet(
        wb.active,
        title="Research Gaps",
        columns=GAP_CHECK_COLUMNS,
        widths=GAP_CHECK_WIDTHS,
        rows=gap_check_rows(gc_results, all_gaps, papers),
        link_columns={"Solving Paper URL"},
    )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gap_check_csv(
    gc_results: list[dict],
    all_gaps: list[dict] | None = None,
    papers: list[dict] | None = None,
) -> str:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=GAP_CHECK_COLUMNS, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(gap_check_rows(gc_results, all_gaps, papers))
    return buffer.getvalue()


